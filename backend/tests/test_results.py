"""Focused Phase 5 results service and endpoint coverage, updated for Part G new export format."""
from __future__ import annotations

import csv
import io
import openpyxl
from datetime import datetime, timedelta

import pytest
from sqlalchemy import create_engine, event
from sqlalchemy.orm import Session

from app.db import Base
from app.models.answer import Answer
from app.models.answer_option import AnswerOption
from app.models.form import Form
from app.models.question import Question
from app.models.question_option import QuestionOption
from app.models.response import Response
from app.models.user import User
from app.services import results_service
from app.services.exceptions import NotFoundError


@pytest.fixture()
def db():
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    with Session(engine) as session:
        user = User(email="results@example.com", name="Results")
        session.add(user)
        session.flush()
        session.add(Form(id=1, user_id=user.id, public_id="results001", title="Results",
                         thank_you_title="Thank you!"))
        session.flush()
        yield session
    Base.metadata.drop_all(engine)


def add_question(db, qid, qtype="short_text", title=None):
    question = Question(id=qid, form_id=1, position=qid, type=qtype, title=title or f"Q{qid}")
    db.add(question)
    db.flush()
    return question


def add_response(db, rid, status="completed", answer=None):
    started = datetime(2025, 1, 1) + timedelta(days=rid)
    submitted = started + timedelta(minutes=5) if status == "completed" else None
    response = Response(id=rid, form_id=1, status=status, started_at=started, submitted_at=submitted)
    db.add(response)
    db.flush()
    if answer is not None:
        db.add(Answer(response_id=rid, question_id=answer[0], value_text=answer[1]))
    db.commit()
    return response


def test_empty_form_list(db):
    assert results_service.list_responses(db, 1, 1).total == 0


def test_empty_form_summary(db):
    summary = results_service.get_summary(db, 1, 1)
    assert summary.total_responses == 0 and summary.completion_rate == 0


def test_empty_form_csv_has_header(db):
    add_question(db, 1)
    first_chunk = next(results_service.export_csv(db, 1, 1))
    # Strip UTF-8 BOM if present, then check new column format starts with '#'
    header = first_chunk.lstrip("\ufeff").lstrip()
    assert header.startswith("#")
    assert "Q1" in header
    assert "Response Type" in header
    assert "Start Date" in header
    assert "Submit Date" in header
    assert "Ending" in header


def test_list_includes_answer_count(db):
    add_question(db, 1)
    add_response(db, 1, answer=(1, "hello"))
    assert results_service.list_responses(db, 1, 1).items[0].answer_count == 1


def test_list_page_one(db):
    for rid in range(1, 4):
        add_response(db, rid)
    assert len(results_service.list_responses(db, 1, 1, 1, 2).items) == 2


def test_list_last_page(db):
    for rid in range(1, 4):
        add_response(db, rid)
    page = results_service.list_responses(db, 1, 1, 2, 2)
    assert len(page.items) == 1 and page.total_pages == 2


def test_list_page_beyond_end(db):
    add_response(db, 1)
    assert results_service.list_responses(db, 1, 1, 99, 10).items == []


def test_list_page_zero_clamps(db):
    add_response(db, 1)
    assert results_service.list_responses(db, 1, 1, 0, 10).page == 1


def test_get_response_scoped_to_form(db):
    add_question(db, 1)
    add_response(db, 1, answer=(1, "hello"))
    response = results_service.get_response(db, 1, 1, 1)
    assert response.answers[0].value_text == "hello"


def test_get_response_missing(db):
    with pytest.raises(NotFoundError):
        results_service.get_response(db, 1, 999, 1)


def test_delete_response(db):
    add_response(db, 1, "completed")
    results_service.delete_response(db, 1, 1, 1)
    with pytest.raises(NotFoundError):
        results_service.get_response(db, 1, 1, 1)


def test_summary_completion_percentage(db):
    add_response(db, 1, "completed")
    add_response(db, 2, "partial")
    summary = results_service.get_summary(db, 1, 1)
    assert summary.completion_rate == 50.0


def test_summary_partial_count(db):
    add_response(db, 1, "partial")
    assert results_service.get_summary(db, 1, 1).partial_responses == 1


def test_summary_choice_zero_counts(db):
    question = add_question(db, 1, "multiple_choice")
    db.add_all([QuestionOption(question_id=question.id, label="A", position=0),
                QuestionOption(question_id=question.id, label="B", position=1)])
    db.commit()
    choices = results_service.get_summary(db, 1, 1).questions[0].choices
    assert [choice["count"] for choice in choices] == [0, 0]


def test_summary_rating_distribution(db):
    add_question(db, 1, "rating")
    for rid, rating in [(1, 3), (2, 3), (3, 5)]:
        response = Response(id=rid, form_id=1, status="completed", started_at=datetime.now())
        db.add(response); db.flush()
        db.add(Answer(response_id=rid, question_id=1, value_number=rating))
    db.commit()
    assert results_service.get_summary(db, 1, 1).questions[0].distribution == {
        "1": 0, "2": 0, "3": 2, "4": 0, "5": 1,
    }


def test_summary_rating_average(db):
    add_question(db, 1, "rating")
    for rid, rating in [(1, 2), (2, 4)]:
        db.add(Response(id=rid, form_id=1, status="completed", started_at=datetime.now()))
        db.flush(); db.add(Answer(response_id=rid, question_id=1, value_number=rating))
    db.commit()
    assert results_service.get_summary(db, 1, 1).questions[0].average == 3.0


def test_summary_number_min_average_max(db):
    add_question(db, 1, "number")
    for rid, number in [(1, 2), (2, 8)]:
        db.add(Response(id=rid, form_id=1, status="completed", started_at=datetime.now()))
        db.flush(); db.add(Answer(response_id=rid, question_id=1, value_number=number))
    db.commit()
    question = results_service.get_summary(db, 1, 1).questions[0]
    assert (question.minimum, question.average, question.maximum) == (2.0, 5.0, 8.0)


def test_summary_latest_text_is_limited(db):
    add_question(db, 1)
    for rid in range(1, 8):
        add_response(db, rid, answer=(1, str(rid)))
    assert len(results_service.get_summary(db, 1, 1).questions[0].latest_answers) == 5


def test_csv_quotes_commas_and_newlines(db):
    add_question(db, 1, title="Question, one")
    add_response(db, 1, answer=(1, "hello,\nworld"))
    content = "".join(results_service.export_csv(db, 1, 1)).lstrip("\ufeff")
    rows = list(csv.reader(io.StringIO(content)))
    # Header: #, Question one, Response Type, Start Date, Submit Date, Ending
    assert rows[0][1] == "Question, one"
    # Data row question column
    assert rows[1][1] == "hello,\nworld"


def test_csv_neutralizes_formula_values(db):
    add_question(db, 1)
    add_response(db, 1, answer=(1, "=SUM(A1:A2)"))
    content = "".join(results_service.export_csv(db, 1, 1))
    assert "'=SUM(A1:A2)" in content


def test_csv_excludes_partial_responses_by_default(db):
    add_question(db, 1)
    add_response(db, 1, "partial")
    content = "".join(results_service.export_csv(db, 1, 1)).lstrip("\ufeff")
    rows = list(csv.reader(io.StringIO(content)))
    # Only header row, no data rows
    assert len(rows) == 1


def test_csv_includes_partial_when_status_partial(db):
    add_question(db, 1)
    add_response(db, 1, "partial")
    content = "".join(results_service.export_csv(db, 1, 1, status="partial")).lstrip("\ufeff")
    rows = list(csv.reader(io.StringIO(content)))
    assert len(rows) == 2
    assert rows[1][2] == "Partial"  # Response Type column


def test_csv_column_order(db):
    """Verify the exact new column layout: #, questions, Response Type, Start Date, Submit Date, Ending."""
    add_question(db, 1, title="My Question")
    add_response(db, 1, answer=(1, "my answer"))
    content = "".join(results_service.export_csv(db, 1, 1)).lstrip("\ufeff")
    rows = list(csv.reader(io.StringIO(content)))
    header = rows[0]
    assert header[0] == "#"
    assert header[1] == "My Question"
    assert header[-4] == "Response Type"
    assert header[-3] == "Start Date"
    assert header[-2] == "Submit Date"
    assert header[-1] == "Ending"


def test_csv_yes_no_text(db):
    """Yes/No questions should output 'Yes' or 'No', not TRUE/FALSE."""
    add_question(db, 1, "yes_no", title="Liked it?")
    started = datetime(2025, 1, 1)
    submitted = started + timedelta(minutes=5)
    response = Response(id=1, form_id=1, status="completed", started_at=started, submitted_at=submitted)
    db.add(response); db.flush()
    db.add(Answer(response_id=1, question_id=1, value_bool=True))
    db.commit()
    content = "".join(results_service.export_csv(db, 1, 1)).lstrip("\ufeff")
    rows = list(csv.reader(io.StringIO(content)))
    assert rows[1][1] == "Yes"


def test_csv_date_format(db):
    """Dates must be 'YYYY-MM-DD HH:MM:SS' format."""
    add_question(db, 1)
    add_response(db, 1, answer=(1, "x"))
    content = "".join(results_service.export_csv(db, 1, 1)).lstrip("\ufeff")
    rows = list(csv.reader(io.StringIO(content)))
    start_date = rows[1][-3]  # Start Date column
    # Must match YYYY-MM-DD HH:MM:SS
    import re
    assert re.match(r"\d{4}-\d{2}-\d{2} \d{2}:\d{2}:\d{2}", start_date)


def test_csv_ending_column_completed(db):
    """Ending column = thank_you_title for completed responses."""
    add_question(db, 1)
    add_response(db, 1, answer=(1, "x"))
    content = "".join(results_service.export_csv(db, 1, 1)).lstrip("\ufeff")
    rows = list(csv.reader(io.StringIO(content)))
    assert rows[1][-1] == "Thank you!"


def test_csv_ids_filter(db):
    """Only selected ids should be exported when ids= is specified."""
    add_question(db, 1)
    add_response(db, 1, answer=(1, "first"))
    add_response(db, 2, answer=(1, "second"))
    content = "".join(results_service.export_csv(db, 1, 1, ids=[1])).lstrip("\ufeff")
    rows = list(csv.reader(io.StringIO(content)))
    assert len(rows) == 2  # header + 1 data row
    assert rows[1][1] == "first"


def test_csv_utfbom(db):
    """CSV must start with a UTF-8 BOM."""
    add_question(db, 1)
    content_bytes = "".join(results_service.export_csv(db, 1, 1)).encode("utf-8")
    assert content_bytes[:3] == b"\xef\xbb\xbf"


def test_csv_streams_completed_rows_without_per_response_queries(db):
    add_question(db, 1)
    for rid in range(1, 4):
        add_response(db, rid, answer=(1, f"answer-{rid}"))

    statements: list[str] = []

    def record_statement(conn, cursor, statement, parameters, context, executemany):
        if statement.lstrip().upper().startswith("SELECT"):
            statements.append(statement)

    engine = db.get_bind()
    event.listen(engine, "before_cursor_execute", record_statement)
    try:
        content = "".join(results_service.export_csv(db, 1, 1))
    finally:
        event.remove(engine, "before_cursor_execute", record_statement)

    assert content.count("answer-") == 3
    # Form lookup, question/header lookup (including its option load), and one
    # streaming export query. There must not be one response query per row.
    assert len(statements) == 4
    assert sum("from responses" in statement.lower() for statement in statements) == 1


def test_csv_joins_multiple_option_labels(db):
    question = add_question(db, 1, "multiple_choice")
    options = [
        QuestionOption(question_id=question.id, label="A", position=0),
        QuestionOption(question_id=question.id, label="B", position=1),
    ]
    db.add_all(options)
    db.flush()
    started = datetime(2025, 1, 2)
    response = Response(id=1, form_id=1, status="completed", started_at=started, submitted_at=started + timedelta(minutes=5))
    db.add(response)
    db.flush()
    answer = Answer(response_id=response.id, question_id=question.id)
    db.add(answer)
    db.flush()
    db.add_all([
        AnswerOption(answer_id=answer.id, option_id=options[0].id),
        AnswerOption(answer_id=answer.id, option_id=options[1].id),
    ])
    db.commit()

    content = "".join(results_service.export_csv(db, 1, 1)).lstrip("\ufeff")
    rows = list(csv.reader(io.StringIO(content)))
    assert rows[1][1] == "A, B"


def test_xlsx_opens_and_header_correct(db):
    """XLSX must open with openpyxl and have correct bold header."""
    add_question(db, 1, title="My Q")
    add_response(db, 1, answer=(1, "hello"))
    data = results_service.export_xlsx(db, 1, 1)
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb.active
    headers = [ws.cell(row=1, column=i).value for i in range(1, ws.max_column + 1)]
    assert headers[0] == "#"
    assert "My Q" in headers
    assert "Response Type" in headers
    assert "Start Date" in headers
    assert "Submit Date" in headers
    assert "Ending" in headers
    # Header row should be bold
    assert ws.cell(row=1, column=1).font.bold is True


def test_xlsx_yes_no_text(db):
    """XLSX Yes/No must be 'Yes'/'No' strings."""
    add_question(db, 1, "yes_no", title="Good?")
    started = datetime(2025, 1, 1)
    response = Response(id=1, form_id=1, status="completed", started_at=started, submitted_at=started + timedelta(minutes=5))
    db.add(response); db.flush()
    db.add(Answer(response_id=1, question_id=1, value_bool=False))
    db.commit()
    data = results_service.export_xlsx(db, 1, 1)
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb.active
    assert ws.cell(row=2, column=2).value == "No"


def test_xlsx_ids_filter(db):
    add_question(db, 1)
    add_response(db, 1, answer=(1, "keep"))
    add_response(db, 2, answer=(1, "drop"))
    data = results_service.export_xlsx(db, 1, 1, ids=[1])
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb.active
    assert ws.max_row == 2  # header + 1 data row
    assert ws.cell(row=2, column=2).value == "keep"


def test_cross_form_isolation(db):
    other = Form(id=2, user_id=1, public_id="results002", title="Other")
    db.add(other); db.commit()
    db.add(Response(id=1, form_id=2, status="completed", started_at=datetime.now())); db.commit()
    assert results_service.list_responses(db, 1, 1).total == 0


def test_cross_user_isolation(db):
    db.add(User(id=2, email="other@example.com", name="Other"))
    db.flush()
    db.add(Form(id=2, user_id=2, public_id="results003", title="Private")); db.commit()
    with pytest.raises(NotFoundError):
        results_service.get_summary(db, 2, 1)
